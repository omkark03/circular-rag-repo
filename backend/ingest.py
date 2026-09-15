"""Document ingestion with upload/read failure detection and OCR fallback.

Checkpoints (any failure is recorded so the dashboard can flag it):

  upload_failed  -> file could not be saved / is empty / wrong type
  read_failed    -> file saved but could not be parsed at all
  empty_text     -> no text layer AND OCR unavailable/failed
  duplicate      -> identical file (SHA-256) already in the repository
  ok             -> indexed and searchable (status_detail notes if OCR was used)

Scanned documents: if a PDF has no extractable text layer, pages are
rasterized with PyMuPDF and read with Tesseract OCR. Image files
(.png/.jpg/.jpeg/.tiff) are OCR'd directly.
"""
import hashlib
import re
from pathlib import Path

import config
import references as refx
import store
import unicodedata

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif"}

# ---------------------------------------------------------------- OCR

def _tesseract_available() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def available_ocr_languages() -> list[str]:
    try:
        import pytesseract
        return sorted(pytesseract.get_languages(config=""))
    except Exception:
        return []


def _detect_skew_angle(gray_img) -> float:
    import cv2
    import numpy as np
    thresh = cv2.threshold(gray_img, 0, 255,
                           cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    coords = np.column_stack(np.where(thresh > 0))
    if len(coords) < 50:
        return 0.0
    angle = cv2.minAreaRect(coords)[-1]
    if angle < -45:
        angle = -(90 + angle)
    else:
        angle = -angle
    return angle


def _preprocess_image(pil_img):
    try:
        import cv2
        import numpy as np
        from PIL import Image

        img = np.array(pil_img.convert("RGB"))
        gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)

        angle = _detect_skew_angle(gray)
        if abs(angle) > 0.1:
            h, w = gray.shape
            m = cv2.getRotationMatrix2D((w // 2, h // 2), angle, 1.0)
            gray = cv2.warpAffine(gray, m, (w, h), flags=cv2.INTER_CUBIC,
                                  borderMode=cv2.BORDER_REPLICATE)

        if config.OCR_DENOISE_ENABLED:
            gray = cv2.fastNlMeansDenoising(gray, None, h=config.OCR_DENOISE_H,
                                            templateWindowSize=7, searchWindowSize=21)
            clahe = cv2.createCLAHE(clipLimit=config.OCR_CLAHE_CLIP, tileGridSize=(8, 8))
            gray = clahe.apply(gray)

        return Image.fromarray(gray)
    except Exception as e:
        print(f"[ingest] OCR pre-processing failed ({type(e).__name__}: {e}) — "
             f"using the original image unprocessed for this page.", flush=True)
        return pil_img


def _ocr_image(pil_img) -> str:
    import pytesseract
    if config.OCR_PREPROCESS:
        pil_img = _preprocess_image(pil_img)
    tess_config = f"--psm {config.OCR_PSM} --oem {config.OCR_OEM}"
    try:
        return pytesseract.image_to_string(pil_img, lang=config.OCR_LANG,
                                           config=tess_config)
    except Exception as e:
        print(f"[ingest] OCR with --psm {config.OCR_PSM} --oem {config.OCR_OEM} "
             f"failed ({type(e).__name__}: {e}) — retrying with Tesseract "
             f"defaults. If this keeps happening, your installed "
             f"'{config.OCR_LANG}' traineddata likely doesn't support "
             f"OCR_OEM={config.OCR_OEM}; try OCR_OEM=3.", flush=True)
        return pytesseract.image_to_string(pil_img, lang=config.OCR_LANG)


def _ocr_pdf(path: Path) -> str:
    if not config.OCR_ENABLED:
        raise RuntimeError(
            "This looks like a scanned PDF (no text layer) and OCR is "
            "disabled. Set OCR_ENABLED=1 to enable it.")
    if not _tesseract_available():
        raise RuntimeError(
            "This looks like a scanned PDF, but Tesseract OCR is not "
            "installed. Install it (Ubuntu: `sudo apt install tesseract-ocr`, "
            "macOS: `brew install tesseract`, Windows: UB-Mannheim installer) "
            "and re-upload.")
    import fitz
    import io
    from PIL import Image

    doc = fitz.open(str(path))
    pages = []
    for i, page in enumerate(doc):
        if i >= config.OCR_MAX_PAGES:
            pages.append(f"[OCR stopped at {config.OCR_MAX_PAGES} pages]")
            break
        pix = page.get_pixmap(dpi=config.OCR_DPI, colorspace=fitz.csRGB, alpha=False)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        pages.append(_ocr_image(img))
    doc.close()
    return "\n\n".join(pages)


def _ocr_image_file(path: Path) -> str:
    if not _tesseract_available():
        raise RuntimeError(
            "Tesseract OCR is not installed — required to read image files. "
            "Install it (Ubuntu: `sudo apt install tesseract-ocr`, macOS: "
            "`brew install tesseract`) and re-upload.")
    from PIL import Image
    with Image.open(str(path)) as img:
        return _ocr_image(img.convert("RGB"))


def _devanagari_ratio(text: str) -> float:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return 0.0
    dev = sum(1 for c in letters if "\u0900" <= c <= "\u097F")
    return dev / len(letters)


def ocr_language_mismatch(text: str) -> bool:
    langs = set(config.OCR_LANG.split("+"))
    if not (langs & {"mar", "hin"}):
        return False
    if len(text.strip()) < 80:
        return False
    if _devanagari_ratio(text) >= config.OCR_MIN_SCRIPT_RATIO:
        return False
    if _looks_like_english_prose(text):
        return False
    return True


_COMMON_ENGLISH_WORDS = {
    "the", "of", "and", "in", "to", "a", "is", "that", "for", "on", "with",
    "as", "this", "by", "or", "be", "are", "was", "were", "not", "has",
    "have", "been", "from", "it", "at", "his", "her", "he", "she", "their",
    "which", "any", "all", "no", "such", "also", "under", "shall", "order",
    "court", "state", "section", "act", "code", "case", "law", "petition",
}


def _looks_like_english_prose(text: str) -> bool:
    words = re.findall(r"[A-Za-z]+", text.lower())
    if len(words) < 30:
        return False
    common = sum(1 for w in words if w in _COMMON_ENGLISH_WORDS)
    return (common / len(words)) > 0.12


# ---------------------------------------------------------------- extraction

def _extract_pdf(path: Path) -> tuple[str, bool]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise RuntimeError(f"PDF is password protected: {e}")
    text = "\n".join((page.extract_text() or "") for page in reader.pages)

    if len(text.strip()) >= config.MIN_TEXT_CHARS:
        return text, False

    ocr_text = _ocr_pdf(path)
    if len(ocr_text.strip()) > len(text.strip()):
        return ocr_text, True
    return text, False


def _extract_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(path: Path) -> tuple[str, bool]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path), False
    if ext == ".txt":
        return path.read_text(errors="replace"), False
    if ext in IMAGE_EXTS:
        return _ocr_image_file(path), True
    raise RuntimeError(f"Unsupported file type: {ext}")


def _slugify_filename(text: str, max_length: int = 80) -> str:
    text = unicodedata.normalize("NFC", text).strip()
    text = re.sub(r'[\\/:*?"<>|\x00-\x1f]', " ", text)
    text = re.sub(r"\s+", "_", text)
    text = text.strip("._")
    return text[:max_length] if text else "circular"


def _rename_stored_file(doc_id: int, old_stored_name: str, title: str,
                        filename: str) -> str:
    """Rename the stored file to a subject-derived name once the title is
    known. Database is updated BEFORE the filesystem rename, and rolled
    back if the rename then fails — this ordering matters: doing it the
    other way round (rename first, then update DB) can leave a file
    physically moved to a new path that the database still doesn't know
    about if the DB update fails for any reason, silently breaking every
    future lookup of this document until manually noticed."""
    ext = Path(filename).suffix or Path(old_stored_name).suffix
    slug = _slugify_filename(title)
    prefix = old_stored_name.split("_", 1)[0]
    new_stored_name = f"{prefix}_{slug}{ext}"
    if new_stored_name == old_stored_name:
        return old_stored_name

    old_path = config.FILES_DIR / old_stored_name
    new_path = config.FILES_DIR / new_stored_name
    counter = 1
    while new_path.exists() and new_stored_name != old_stored_name:
        new_stored_name = f"{prefix}_{slug}_{counter}{ext}"
        new_path = config.FILES_DIR / new_stored_name
        counter += 1

    if not old_path.exists():
        return old_stored_name

    try:
        store.update_document(doc_id, stored_name=new_stored_name)
    except Exception as e:
        print(f"[ingest] Could not update database for doc {doc_id} rename "
             f"({type(e).__name__}: {e}) — keeping original filename, "
             f"file left untouched.", flush=True)
        return old_stored_name

    try:
        old_path.rename(new_path)
        return new_stored_name
    except Exception as e:
        print(f"[ingest] Database updated but file rename failed for doc "
             f"{doc_id} ({type(e).__name__}: {e}) — rolling back database "
             f"to keep them in sync.", flush=True)
        try:
            store.update_document(doc_id, stored_name=old_stored_name)
        except Exception:
            print(f"[ingest] CRITICAL: could not roll back stored_name for "
                 f"doc {doc_id} after a failed file rename — database and "
                 f"filesystem may now be out of sync. Manually verify "
                 f"stored_name for this document.", flush=True)
        return old_stored_name


# ---------------------------------------------------------------- helpers

def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _is_running_header(line: str, full_text: str) -> bool:
    if len(line) < 6:
        return True
    return full_text.count(line) >= 3


def _guess_metadata(text: str, filename: str) -> dict:
    head_lines = [l.strip() for l in text[:2000].splitlines() if l.strip()]
    issuer_candidates = [l for l in head_lines if not _is_running_header(l, text)]
    pool = issuer_candidates or head_lines
    issuer = pool[0][:120] if pool else "Unknown issuer"
    title = None
    for line in head_lines[1:8]:
        if re.search(r"(sub(ject)?\s*[:\-])", line, re.IGNORECASE):
            title = re.sub(r"^sub(ject)?\s*[:\-]\s*", "", line, flags=re.IGNORECASE)
            break
    if not title:
        title = Path(filename).stem.replace("_", " ").replace("-", " ").title()
    m = re.search(r"\b(19|20)\d{2}\b", text[:2500]) or re.search(r"\b(19|20)\d{2}\b", filename)
    year = m.group(0) if m else "n.d."
    return {"issuer": issuer, "title": title[:200], "year": year}


def chunk_text(text: str) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text)
    chunks, step = [], config.CHUNK_SIZE - config.CHUNK_OVERLAP
    for i in range(0, max(len(text), 1), step):
        piece = text[i : i + config.CHUNK_SIZE].strip()
        if piece:
            chunks.append(piece)
        if i + config.CHUNK_SIZE >= len(text):
            break
    return chunks


# ---------------------------------------------------------------- vector store

_embedder = None
_collection = None


def get_collection():
    global _collection
    if _collection is None:
        import chromadb
        client = chromadb.PersistentClient(path=config.CHROMA_DIR)
        _collection = client.get_or_create_collection(
            "circulars", metadata={"hnsw:space": "cosine"})
    return _collection


def get_embedder():
    global _embedder
    if _embedder is None:
        from sentence_transformers import SentenceTransformer
        _embedder = SentenceTransformer(config.EMBED_MODEL, device=config.EMBED_DEVICE)
    return _embedder


def _embedding_text(chunk: str, meta: dict) -> str:
    header_parts = [str(v) for k in ("title", "issuer", "year")
                    if (v := meta.get(k))]
    header = " | ".join(header_parts)
    return f"{header}\n{chunk}" if header else chunk


def index_chunks(doc_id: int, meta: dict, chunks: list[str]):
    col = get_collection()
    embedding_texts = [_embedding_text(c, meta) for c in chunks]
    emb = get_embedder().encode(embedding_texts, show_progress_bar=False).tolist()
    col.add(
        ids=[f"{doc_id}:{i}" for i in range(len(chunks))],
        embeddings=emb,
        documents=chunks,
        metadatas=[{**meta, "doc_id": doc_id, "chunk": i} for i in range(len(chunks))],
    )


# ---------------------------------------------------------------- main entry

def finalize_document(doc_id: int, filename: str, stored_name: str,
                      stored_path: Path) -> dict:
    ocr_used = False
    try:
        text, ocr_used = extract_text(stored_path)
    except RuntimeError as e:
        msg = str(e)
        status = "empty_text" if ("OCR" in msg or "Tesseract" in msg) else "read_failed"
        store.update_document(doc_id, status=status, status_detail=msg, n_chunks=0)
        return {"status": status, "detail": msg}
    except Exception as e:
        msg = f"Text extraction failed: {e}"
        store.update_document(doc_id, status="read_failed", status_detail=msg, n_chunks=0)
        return {"status": "read_failed", "detail": msg}

    if ocr_used and ocr_language_mismatch(text):
        ratio = _devanagari_ratio(text)
        detail = (
            f"OCR ran (OCR_LANG={config.OCR_LANG}) but only {ratio:.0%} of letters in "
            f"the output are Devanagari (threshold: {config.OCR_MIN_SCRIPT_RATIO:.0%}). "
            "This usually means the 'mar' language pack isn't installed.")
        store.update_document(doc_id, status="ocr_suspect", status_detail=detail,
                              n_chunks=0, circular_no=None, text_sample=text[:500])
        return {"status": "ocr_suspect", "detail": detail, "ocr": True}

    if len(text.strip()) < config.MIN_TEXT_CHARS:
        detail = ("File was OCR'd but produced almost no readable text."
                  if ocr_used else
                  "File opened but contains almost no extractable text.")
        store.update_document(doc_id, status="empty_text", status_detail=detail,
                              n_chunks=0, text_sample=text[:500])
        return {"status": "empty_text", "detail": detail}

    # -- index -----------------------------------------------------------
    meta = _guess_metadata(text, filename)
    own_no = refx.detect_own_number(text, filename)
    stored_name = _rename_stored_file(doc_id, stored_name, meta["title"], filename)
    store.update_document(
        doc_id, circular_no=own_no, status="ok",
        status_detail="Indexed via OCR (scanned document)" if ocr_used else "Indexed",
        text_sample=text[:500], **meta)
    chunks = chunk_text(text)
    try:
        get_collection().delete(where={"doc_id": doc_id})
    except Exception:
        pass
    try:
        index_chunks(doc_id, {"filename": filename, "stored_name": stored_name, **meta},
                     chunks)
        store.update_document(doc_id, n_chunks=len(chunks))
    except Exception as e:
        msg = f"Indexing failed: {e}"
        store.update_document(doc_id, status="read_failed", status_detail=msg)
        return {"status": "read_failed", "detail": msg}

    store.clear_references(doc_id)
    if store.cross_reference_check_enabled():
        for raw, norm in refx.split_body_references(text, own_no):
            store.add_reference(doc_id, raw, norm)
        store.resolve_pending_references()

    return {"status": "ok", "circular_no": own_no, "ocr": ocr_used,
            "detail": f"Indexed {len(chunks)} chunks"
                      + (" · read via OCR (scanned document)" if ocr_used else "") + "."}


def ingest_upload(filename: str, data: bytes) -> dict:
    store.init_db()
    result = {"filename": filename, "status": "ok", "detail": ""}

    ext = Path(filename).suffix.lower()
    if not data:
        doc_id = store.add_document(filename=filename, status="upload_failed",
                                    status_detail="Uploaded file is empty (0 bytes).")
        return {**result, "status": "upload_failed",
                "detail": "Uploaded file is empty (0 bytes).", "doc_id": doc_id}
    if ext not in config.ALLOWED_EXTENSIONS:
        doc_id = store.add_document(filename=filename, status="upload_failed",
                                    status_detail=f"Unsupported file type '{ext}'.")
        return {**result, "status": "upload_failed",
                "detail": f"Unsupported file type '{ext}'.", "doc_id": doc_id}

    sha = _sha256(data)
    dup = store.find_by_hash(sha)
    if dup:
        return {**result, "status": "duplicate",
                "detail": f"Not uploaded — identical to '{dup['filename']}' "
                          f"(already in the repository).",
                "duplicate_of": dup["id"]}

    stored_name = f"{sha[:10]}_{re.sub(r'[^A-Za-z0-9._-]', '_', filename)}"
    stored_path = config.FILES_DIR / stored_name
    try:
        stored_path.write_bytes(data)
    except OSError as e:
        doc_id = store.add_document(filename=filename, status="upload_failed",
                                    status_detail=f"Could not save file: {e}")
        return {**result, "status": "upload_failed", "detail": str(e), "doc_id": doc_id}

    doc_id = store.add_document(filename=filename, stored_name=stored_name, sha256=sha,
                                status="read_failed", status_detail="Processing…")
    outcome = finalize_document(doc_id, filename, stored_name, stored_path)
    result.update(outcome)
    result["doc_id"] = doc_id
    return result


# ---------------------------------------------------------------- deletion

def delete_document(doc_id: int) -> dict | None:
    row = store.delete_document(doc_id)
    if not row:
        return None
    try:
        get_collection().delete(where={"doc_id": doc_id})
    except Exception:
        pass
    if row.get("stored_name"):
        try:
            (config.FILES_DIR / row["stored_name"]).unlink(missing_ok=True)
        except OSError:
            pass
    return row
