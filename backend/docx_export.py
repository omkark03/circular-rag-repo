"""Minimal Markdown -> DOCX converter for exporting generated templates.

Deliberately lightweight (no markdown-it/mistune dependency): handles the
subset of Markdown the LLM actually produces for templates — headings,
bold/italic inline emphasis, bullet/numbered lists, and GFM tables.
"""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

_BOLD_ITALIC = re.compile(r"\*\*\*(.+?)\*\*\*")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")


def _add_inline(paragraph, text: str):
    """Add text to a docx paragraph, honoring **bold**, *italic*, ***both***."""
    tokens = []  # list of (text, bold, italic)
    pos = 0
    pattern = re.compile(r"\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*")
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            tokens.append((m.group(1), True, True))
        elif m.group(2) is not None:
            tokens.append((m.group(2), True, False))
        else:
            tokens.append((m.group(3), False, True))
        pos = m.end()
    if pos < len(text):
        tokens.append((text[pos:], False, False))
    if not tokens:
        paragraph.add_run(text)
        return
    for t, bold, italic in tokens:
        run = paragraph.add_run(t)
        run.bold = bold
        run.italic = italic


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def _is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", line.strip()))


def _split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color})
    tcPr.append(shd)


def markdown_to_docx(markdown_text: str, title: str = "Template") -> bytes:
    """Render Markdown to a .docx file, returned as bytes."""
    doc = Document()
    doc.core_properties.title = title

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = markdown_text.replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # ---- table block ----
        if _is_table_row(stripped):
            rows = []
            while i < n and _is_table_row(lines[i].strip()):
                rows.append(lines[i].strip())
                i += 1
            if len(rows) >= 2 and _is_separator_row(rows[1]):
                header_cells = _split_row(rows[0])
                data_rows = [_split_row(r) for r in rows[2:]]
            else:
                header_cells = _split_row(rows[0])
                data_rows = [_split_row(r) for r in rows[1:]]
            ncols = len(header_cells)
            table = doc.add_table(rows=1, cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for ci, text in enumerate(header_cells):
                cell = table.rows[0].cells[ci]
                cell.paragraphs[0].clear()
                _add_inline(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                _shade_cell(cell, "EDEDE6")
            for row_cells in data_rows:
                row = table.add_row()
                for ci in range(ncols):
                    text = row_cells[ci] if ci < len(row_cells) else ""
                    row.cells[ci].paragraphs[0].clear()
                    _add_inline(row.cells[ci].paragraphs[0], text)
            doc.add_paragraph()  # spacing after table
            continue

        # ---- headings ----
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            p = doc.add_heading(level=level)
            _add_inline(p, m.group(2))
            i += 1
            continue

        # ---- bullet list ----
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1))
            i += 1
            continue

        # ---- numbered list ----
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(1))
            i += 1
            continue

        # ---- blockquote ----
        m = re.match(r"^>\s?(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_inline(p, m.group(1))
            i += 1
            continue

        # ---- plain paragraph ----
        p = doc.add_paragraph()
        _add_inline(p, stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
