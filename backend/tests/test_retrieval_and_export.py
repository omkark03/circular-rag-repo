"""Regression tests for retrieval.py (relevance filtering) and
docx_export.py (Markdown -> Word conversion for generated templates).

Run with:  pytest tests/test_retrieval_and_export.py -v
Or standalone:  python tests/test_retrieval_and_export.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import config
import retrieval as rt


# ---------------------------------------------------------------- relevance filtering

def test_irrelevant_document_excluded_alongside_relevant_one():
    """The core fix: a genuinely relevant chunk plus an unrelated one (that
    only surfaced via keyword overlap) must not both survive — only the
    relevant one should."""
    candidates = [
        {"meta": {"doc_id": 1}, "rerank_score": 6.2},
        {"meta": {"doc_id": 2}, "rerank_score": -7.5},
    ]
    kept = rt._filter_relevant(candidates, "rerank_score",
                               config.RERANK_MIN_SCORE, config.RERANK_MAX_GAP)
    assert [c["meta"]["doc_id"] for c in kept] == [1]


def test_nothing_relevant_returns_empty():
    """When nothing clears the bar, the app should honestly report nothing
    found rather than force-citing the closest-available junk."""
    candidates = [
        {"meta": {"doc_id": 1}, "rerank_score": -6.0},
        {"meta": {"doc_id": 2}, "rerank_score": -6.5},
        {"meta": {"doc_id": 3}, "rerank_score": -5.8},
    ]
    kept = rt._filter_relevant(candidates, "rerank_score",
                               config.RERANK_MIN_SCORE, config.RERANK_MAX_GAP)
    assert kept == []


def test_multiple_genuinely_relevant_all_survive():
    candidates = [
        {"meta": {"doc_id": 1}, "rerank_score": 5.0},
        {"meta": {"doc_id": 2}, "rerank_score": 4.2},
        {"meta": {"doc_id": 3}, "rerank_score": 3.9},
    ]
    kept = rt._filter_relevant(candidates, "rerank_score",
                               config.RERANK_MIN_SCORE, config.RERANK_MAX_GAP)
    assert sorted(c["meta"]["doc_id"] for c in kept) == [1, 2, 3]


def test_absolute_floor_boundary():
    candidates = [
        {"meta": {"doc_id": 1}, "rerank_score": config.RERANK_MIN_SCORE + 0.1},
        {"meta": {"doc_id": 2}, "rerank_score": config.RERANK_MIN_SCORE - 0.1},
    ]
    kept = rt._filter_relevant(candidates, "rerank_score",
                               config.RERANK_MIN_SCORE, config.RERANK_MAX_GAP)
    assert [c["meta"]["doc_id"] for c in kept] == [1]


def test_end_to_end_retrieve_excludes_irrelevant_document(monkeypatch):
    """Full retrieve() pipeline: vector search returns a relevant and an
    irrelevant chunk; reranker scores confirm the irrelevant one should be
    dropped from the final result."""
    monkeypatch.setattr(rt, "_vector_candidates", lambda q, k: [
        {"text": "KYC norms for savings accounts...",
         "meta": {"doc_id": 1, "chunk": 0}, "vec_score": 0.81},
        {"text": "Traffic signal maintenance schedule...",
         "meta": {"doc_id": 2, "chunk": 0}, "vec_score": 0.34},
    ])
    monkeypatch.setattr(rt, "_bm25_candidates", lambda q, k: [])
    monkeypatch.setattr(rt, "_get_reranker", lambda: "mock-active")

    def fake_rerank(question, candidates, top_k):
        scores = {1: 6.8, 2: -8.2}
        for c in candidates:
            c["rerank_score"] = scores[c["meta"]["doc_id"]]
        return sorted(candidates, key=lambda c: c["rerank_score"], reverse=True)[:top_k]
    monkeypatch.setattr(rt, "_rerank", fake_rerank)

    hits = rt.retrieve("What is the KYC limit for small accounts?")
    assert [h["meta"]["doc_id"] for h in hits] == [1]


# ---------------------------------------------------------------- docx export

# ---------------------------------------------------------------- leaked-marker detection

def test_leak_detector_silent_on_clean_answer():
    import rag
    clean = "The Superintendent of Police must inspect each station (Maharashtra Police, 1966)."
    assert rag._check_for_leaked_markers(clean) == []


def test_leak_detector_catches_placeholder_text():
    """Regression: this exact failure (the model parroting literal example/
    instruction text as if it were a real citation) recurred three times
    with three different literal strings used as examples -- this is a
    visibility net so it's caught immediately if it happens a fourth time,
    since prompt wording alone proved unreliable against it."""
    import rag
    leaked = "This is required (ISSUER-PLACEHOLDER, 19XX) according to the rules."
    found = rag._check_for_leaked_markers(leaked)
    assert "ISSUER-PLACEHOLDER" in found
    assert "19XX" in found


def test_prompts_contain_no_literal_example_citation_text():
    """The prompts themselves must never contain a concrete example
    citation string -- that's exactly what got parroted before. Verifying
    the instruction relies on describing the mechanism, not demonstrating
    it with copyable sample text."""
    import rag
    for marker in ("ISSUER-PLACEHOLDER", "19XX", "Mumbai Police", "Reserve Bank of India"):
        assert marker not in rag.SYSTEM_PROMPT, f"{marker} found in SYSTEM_PROMPT"
        assert marker not in rag.TEMPLATE_SYSTEM_PROMPT, f"{marker} found in TEMPLATE_SYSTEM_PROMPT"


def test_markdown_to_docx_preserves_structure():
    import docx_export
    from docx import Document

    md = """# Panchnama

**Officer:** ____________________

## Items

| Sr. No. | Description | Quantity |
|---|---|---|
| 1 | ___________ | ___ |
| 2 | ___________ | ___ |

- Witness 1: ____________________
- Witness 2: ____________________
"""
    data = docx_export.markdown_to_docx(md, title="Test Template")
    assert len(data) > 500

    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Panchnama" in h for h in headings)
    assert any("Items" in h for h in headings)

    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 2

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold  # header bold

    bold_found = any(r.bold for p in doc.paragraphs for r in p.runs
                     if "Officer" in p.text)
    assert bold_found


if __name__ == "__main__":
    import unittest.mock as mock

    class _FakeMonkeypatch:
        def __init__(self):
            self._patches = []

        def setattr(self, obj, name, value):
            original = getattr(obj, name)
            self._patches.append((obj, name, original))
            setattr(obj, name, value)

        def undo(self):
            for obj, name, original in reversed(self._patches):
                setattr(obj, name, original)

    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            if "monkeypatch" in fn.__code__.co_varnames[:fn.__code__.co_argcount]:
                mp = _FakeMonkeypatch()
                try:
                    fn(mp)
                finally:
                    mp.undo()
            else:
                fn()
            print(f"PASS  {name}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {name}: {e}")
        except Exception as e:
            failed += 1
            print(f"ERROR {name}: {type(e).__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
